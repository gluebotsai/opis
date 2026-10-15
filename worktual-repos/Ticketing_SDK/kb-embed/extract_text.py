"""Extract exact upload text. No rewrite, no invented copy."""

from __future__ import annotations

import io
import re
from typing import Callable

from extract_pdf import extract_pdf_text

SUPPORTED_TEXT = {".txt", ".md", ".csv"}
SUPPORTED_RICH = {".pdf", ".docx", ".rtf"}
SUPPORTED_IMAGE = {".png", ".jpg", ".jpeg"}
UNSUPPORTED = {".doc", ".ppt", ".pptx", ".xls", ".xlsx"}

EXTRACTABLE = SUPPORTED_TEXT | SUPPORTED_RICH | SUPPORTED_IMAGE


class ExtractError(ValueError):
    """User-facing extract failure (unsupported type or missing OCR)."""


def file_suffix(filename: str) -> str:
    name = str(filename or "").strip().lower()
    if "." not in name:
        return ""
    return "." + name.rsplit(".", 1)[-1]


def extract_file_text(file_bytes: bytes, filename: str) -> str:
    data = file_bytes or b""
    if not data:
        return ""
    suffix = file_suffix(filename)
    if suffix in UNSUPPORTED:
        raise ExtractError(
            "This file type cannot be extracted. Convert it or paste the text."
        )
    handlers: dict[str, Callable[[bytes], str]] = {
        ".pdf": extract_pdf_text,
        ".docx": _extract_docx,
        ".rtf": _extract_rtf,
        ".txt": _extract_plain,
        ".md": _extract_plain,
        ".csv": _extract_plain,
        ".png": _extract_image,
        ".jpg": _extract_image,
        ".jpeg": _extract_image,
    }
    handler = handlers.get(suffix)
    if handler is None:
        raise ExtractError(
            "Upload a PDF, DOCX, TXT, MD, CSV, RTF, or image file."
        )
    return handler(data).strip()


def _extract_plain(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="ignore").strip()


def _extract_docx(file_bytes: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(file_bytes))
    parts: list[str] = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_rtf(file_bytes: bytes) -> str:
    raw = file_bytes.decode("latin-1", errors="ignore")
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+\d* ?", " ", raw)
    raw = re.sub(r"[{}]", " ", raw)
    raw = re.sub(r"\s+", " ", raw)
    return raw.strip()


def _extract_image(file_bytes: bytes) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise ExtractError(
            "Image OCR is not available. Paste the article text instead."
        ) from exc
    img = Image.open(io.BytesIO(file_bytes))
    return (pytesseract.image_to_string(img) or "").strip()
